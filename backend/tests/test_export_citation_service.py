from __future__ import annotations

import io
import zipfile

from docx import Document

from app.services.export_citation_service import (
    add_external_hyperlink,
    source_trace,
    source_trace_lines,
)


def _hit() -> dict:
    return {
        "id": "source-1",
        "urn": "urn:lex:br:federal:lei:2024-01-01;123",
        "pinpoint": "Art. 12, § 1º",
        "validity": "vigente",
        "status_as_of": "2026-07-17",
        "review_status": "provisional",
        "verification_scope": "official metadata and local excerpt hash",
        "last_verified_at": "2026-07-17T08:30:00Z",
        "official_url": "https://www.planalto.gov.br/example",
    }


def test_source_trace_keeps_full_audit_envelope() -> None:
    trace = source_trace(_hit())

    assert trace == {
        "citation_id": "urn:lex:br:federal:lei:2024-01-01;123",
        "pinpoint": "Art. 12, § 1º",
        "validity": "vigente",
        "status_as_of": "2026-07-17",
        "review_status": "provisional",
        "verification_scope": "official metadata and local excerpt hash",
        "last_verified_at": "2026-07-17T08:30:00Z",
        "official_url": "https://www.planalto.gov.br/example",
    }
    text = "\n".join(source_trace_lines(_hit()))
    assert "Art. 12" in text
    assert "2026-07-17" in text
    assert "provisional" in text
    assert "official metadata" in text


def test_source_trace_prefers_explicit_cross_jurisdiction_citation_id() -> None:
    hit = _hit()
    hit["citation_id"] = "mx-federal-do:2026-07-17:42"

    assert source_trace(hit)["citation_id"] == "mx-federal-do:2026-07-17:42"


def test_docx_hyperlink_is_external_and_clickable() -> None:
    doc = Document()
    paragraph = doc.add_paragraph("官方链接：")
    add_external_hyperlink(
        paragraph,
        url="https://www.planalto.gov.br/example",
        label="官方原文",
    )
    output = io.BytesIO()
    doc.save(output)

    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        relationships = archive.read("word/_rels/document.xml.rels").decode("utf-8")

    assert "w:hyperlink" in document_xml
    assert "官方原文" in document_xml
    assert "TargetMode=\"External\"" in relationships
    assert "https://www.planalto.gov.br/example" in relationships


def test_missing_source_fields_are_not_silently_omitted() -> None:
    trace = source_trace({"id": "local-record"})

    assert trace["citation_id"] == "local-record"
    assert trace["pinpoint"] == "—"
    assert trace["status_as_of"] == "—"
    assert trace["review_status"] == "pending"
    assert trace["official_url"] == ""


def test_docx_hyperlink_rejects_non_https_relationships() -> None:
    doc = Document()
    paragraph = doc.add_paragraph("官方链接：")
    add_external_hyperlink(
        paragraph,
        url="file:///Users/reviewer/private.txt",
        label="打开本地文件",
    )
    output = io.BytesIO()
    doc.save(output)

    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        relationships = archive.read("word/_rels/document.xml.rels").decode("utf-8")

    assert "仅允许 HTTPS 官方来源" in document_xml
    assert "file:///" not in relationships
